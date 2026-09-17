from collections import defaultdict
from pathlib import Path

from docx import Document

from app.config import BACKEND_ROOT
from app.db.models import WordDocument


def resolve_word_document_path(document: WordDocument) -> Path:
    if document.file_path:
        path = Path(document.file_path)
        if path.exists():
            return path

    upload_dir = BACKEND_ROOT / "data" / "uploads" / "word"
    for path in (
        upload_dir / f"{document.id}_{document.file_name}",
        upload_dir / document.file_name,
    ):
        if path.exists():
            return path
    raise FileNotFoundError(f"找不到 Word 文件：{document.file_name}")


def apply_word_changes(
    source_path: Path, output_path: Path, changes: list[dict[str, object]]
) -> list[dict[str, object]]:
    document = Document(str(source_path))
    grouped: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    for change in changes:
        change["status"] = "pending"
        grouped[locator_key(dict(change["locator"]))].append(change)

    for paragraph_changes in grouped.values():
        paragraph_changes.sort(
            key=lambda item: int(dict(item["span"])["start"]), reverse=True
        )
        for change in paragraph_changes:
            try:
                paragraph = locate_paragraph(document, dict(change["locator"]))
                replace_paragraph_span(
                    paragraph,
                    span=dict(change["span"]),
                    replacement=str(change["after"]),
                )
                change["status"] = "applied"
            except (IndexError, KeyError, ValueError) as exc:
                change["status"] = "failed"
                change["error"] = str(exc)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return changes


def validate_word_change_diff(
    source_path: Path,
    output_path: Path,
    changes: list[dict[str, object]],
) -> dict[str, object]:
    """Check that a Word copy contains exactly the planned text replacements."""
    source_text = document_text_by_locator(Document(str(source_path)))
    output_text = document_text_by_locator(Document(str(output_path)))
    planned: dict[tuple[object, ...], list[dict[str, object]]] = defaultdict(list)
    errors: list[str] = []

    for change in changes:
        before = str(change.get("before") or "")
        after = str(change.get("after") or "")
        if not before or before == after:
            errors.append(f"change_not_changed:{change.get('change_id', '')}")
            continue
        try:
            planned[locator_key(dict(change.get("locator") or {}))].append(change)
        except (KeyError, TypeError, ValueError):
            errors.append(f"invalid_locator:{change.get('change_id', '')}")

    verified_change_count = 0
    for key, batch in planned.items():
        original = source_text.get(key)
        if original is None:
            errors.append(f"source_locator_missing:{key}")
            continue
        expected = original
        for change in sorted(
            batch,
            key=lambda item: int(dict(item.get("span") or {}).get("start", -1)),
            reverse=True,
        ):
            span = dict(change.get("span") or {})
            try:
                start = int(span["start"])
                end = int(span["end"])
            except (KeyError, TypeError, ValueError):
                errors.append(f"invalid_span:{change.get('change_id', '')}")
                continue
            before = str(change.get("before") or "")
            if (
                start < 0
                or end <= start
                or end > len(expected)
                or expected[start:end] != before
            ):
                errors.append(f"source_text_mismatch:{change.get('change_id', '')}")
                continue
            expected = (
                expected[:start] + str(change.get("after") or "") + expected[end:]
            )
            verified_change_count += 1
        if output_text.get(key) != expected:
            errors.append(f"target_text_mismatch:{key}")

    changed_locators = {
        key
        for key in set(source_text) | set(output_text)
        if source_text.get(key) != output_text.get(key)
    }
    unexpected_locators = sorted(changed_locators - set(planned), key=str)
    if unexpected_locators:
        errors.append("outside_target_change")

    return {
        "ok": not errors and verified_change_count == len(changes),
        "errors": errors,
        "planned_change_count": len(changes),
        "verified_change_count": verified_change_count,
        "changed_locator_count": len(changed_locators),
        "unexpected_locators": [list(item) for item in unexpected_locators],
    }


def document_text_by_locator(document) -> dict[tuple[object, ...], str]:
    values = {
        ("body_paragraph", index): paragraph.text
        for index, paragraph in enumerate(document.paragraphs)
    }
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    values[
                        (
                            "table_cell_paragraph",
                            table_index,
                            row_index,
                            cell_index,
                            paragraph_index,
                        )
                    ] = paragraph.text
    return values


def locator_key(locator: dict[str, object]) -> tuple[object, ...]:
    if locator.get("container") == "body_paragraph":
        return ("body_paragraph", int(locator["paragraph_index"]))
    return (
        "table_cell_paragraph",
        int(locator["table_index"]),
        int(locator["row_index"]),
        int(locator["cell_index"]),
        int(locator["paragraph_index"]),
    )


def locate_paragraph(document, locator: dict[str, object]):
    if locator.get("container") == "body_paragraph":
        return document.paragraphs[int(locator["paragraph_index"])]
    if locator.get("container") == "table_cell_paragraph":
        table = document.tables[int(locator["table_index"])]
        cell = table.rows[int(locator["row_index"])].cells[int(locator["cell_index"])]
        return cell.paragraphs[int(locator["paragraph_index"])]
    raise ValueError("引用缺少可写入的 Word 容器位置")


def replace_paragraph_span(
    paragraph, *, span: dict[str, object], replacement: str
) -> None:
    start = int(span["start"])
    end = int(span["end"])
    expected = str(span.get("text") or "")
    text = paragraph.text
    if start < 0 or end <= start or end > len(text):
        raise ValueError("引用字符区间已失效")
    if expected and text[start:end] != expected:
        raise ValueError(f"定位文本不一致，预期“{expected}”，实际“{text[start:end]}”")

    runs = paragraph_runs(paragraph)
    run_ranges: list[tuple[object, int, int]] = []
    cursor = 0
    for run in runs:
        run_end = cursor + len(run.text)
        run_ranges.append((run, cursor, run_end))
        cursor = run_end
    if cursor != len(text):
        raise ValueError("该段落包含暂不支持的 Word 域或特殊对象")

    first_index = next(
        (
            index
            for index, (_, left, right) in enumerate(run_ranges)
            if left <= start < right
        ),
        None,
    )
    last_index = next(
        (
            index
            for index, (_, left, right) in enumerate(run_ranges)
            if left < end <= right
        ),
        None,
    )
    if first_index is None or last_index is None:
        raise ValueError("无法将引用字符区间映射到 Word 文本样式")

    first_run, first_left, _ = run_ranges[first_index]
    last_run, last_left, _ = run_ranges[last_index]
    if first_index == last_index:
        first_run.text = (
            first_run.text[: start - first_left]
            + replacement
            + first_run.text[end - first_left :]
        )
        return

    first_run.text = first_run.text[: start - first_left] + replacement
    for index in range(first_index + 1, last_index):
        run_ranges[index][0].text = ""
    last_run.text = last_run.text[end - last_left :]


def paragraph_runs(paragraph) -> list[object]:
    runs: list[object] = []
    for item in paragraph.iter_inner_content():
        if hasattr(item, "runs"):
            runs.extend(item.runs)
        else:
            runs.append(item)
    return runs
