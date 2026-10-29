from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.agent import tools
from app.db.models import Base, User, WordDocument, WordReference
from app.word_loader import paragraph_style_name, parse_docx_file


@compiles(JSONB, "sqlite")
def compile_jsonb_for_test(element, compiler, **kwargs):
    del element, compiler, kwargs
    return "JSON"


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    document = Document()
    document.add_heading("1 概述", level=1)
    document.add_paragraph("本工程依据 GB 50311-2016 第 3.1.1 条。")
    document.add_heading("1.1 设计依据", level=2)
    document.add_paragraph("正文", style="List Bullet")
    plain = document.add_paragraph("缺失样式的段落")
    plain._p.get_or_add_pPr().get_or_add_pStyle().val = "NoSuchStyle"
    table = document.add_table(rows=3, cols=3)
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.text = f"r{row_index}c{cell_index}"
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(1, 1).merge(table.cell(2, 1))
    path = tmp_path / "sample.docx"
    document.save(path)
    return path


def test_cached_style_name_matches_python_docx(sample_docx):
    document = Document(sample_docx)
    names = [paragraph_style_name(p) for p in document.paragraphs]
    assert names == [p.style.name for p in document.paragraphs]
    assert "Heading 1" in names and "List Bullet" in names
    # Unknown style ids fall back to the default paragraph style.
    assert names[4] == "Normal"


def test_merged_cells_are_kept_once_and_deterministically(sample_docx):
    runs = [
        [chunk.text for chunk in parse_docx_file(str(sample_docx)) if chunk.source_type == "word_docx_table"]
        for _ in range(20)
    ]
    assert all(run == runs[0] for run in runs)
    # Merging moves the cells' paragraphs into one cell: every text appears exactly once.
    assert sorted(runs[0]) == sorted(f"r{r}c{c}" for r in range(3) for c in range(3))


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    open_sessions = []

    class TrackingSession(Session):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            open_sessions.append(self)

        def close(self):
            if self in open_sessions:
                open_sessions.remove(self)
            super().close()

    factory = sessionmaker(bind=engine, class_=TrackingSession)
    with factory() as session:
        session.add(User(id=7, username="visitor", password_hash=""))
        session.commit()
    with patch.object(tools, "SessionLocal", factory):
        yield factory, open_sessions


def test_parse_runs_without_holding_a_db_session(db, sample_docx, tmp_path):
    factory, open_sessions = db
    sessions_during_parse = []
    real_parse = tools.parse_docx_file

    def parse_and_record(path):
        sessions_during_parse.append(len(open_sessions))
        return real_parse(path)

    with (
        patch.object(tools, "BACKEND_ROOT", tmp_path),
        patch.object(tools, "UPLOAD_ROOT", tmp_path),
        patch.object(tools, "parse_docx_file", side_effect=parse_and_record),
    ):
        result = tools.parse_word_tool(
            file_path=str(sample_docx), owner_user_id=7, file_name="设计说明.docx"
        )

    assert sessions_during_parse == [0]
    assert result["parse_status"] == "parsed"
    with factory() as session:
        document = session.get(WordDocument, result["document_id"])
        assert document.file_name == "设计说明.docx"
        assert Path(document.file_path).name == f"{document.id}_设计说明.docx"
        assert session.query(WordReference).filter_by(document_id=document.id).count() >= 1


def test_failed_parse_marks_document_failed(db, tmp_path):
    factory, _ = db
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a docx")

    with patch.object(tools, "UPLOAD_ROOT", tmp_path), pytest.raises(Exception):
        tools.parse_word_tool(file_path=str(broken), owner_user_id=7)

    with factory() as session:
        document = session.query(WordDocument).one()
        assert document.parse_status == "failed"
        assert document.file_name == "broken.docx"
