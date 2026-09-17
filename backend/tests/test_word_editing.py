import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.services.word_editor import apply_word_changes, validate_word_change_diff
from app.services.word_references import extract_references
from app.word_loader import parse_docx_file


class WordEditingTest(unittest.TestCase):
    def test_extracts_exact_spans_and_preserves_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            output = Path(directory) / "output.docx"
            document = Document()
            document.add_heading("1.2 设计依据", level=1)
            paragraph = document.add_paragraph()
            paragraph.add_run("依据《旧").bold = True
            paragraph.add_run("标准》GB 100-2010 执行。")
            document.save(source)

            source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            reference = extract_references(parse_docx_file(str(source)))[0]
            locator = dict(reference["locator"])
            changes = [
                make_change(locator, "standard_name", "新标准"),
                make_change(locator, "standard_no", "GB 100-2026"),
            ]

            result = apply_word_changes(source, output, changes)

            self.assertTrue(all(change["status"] == "applied" for change in result))
            self.assertEqual(
                source_hash, hashlib.sha256(source.read_bytes()).hexdigest()
            )
            self.assertEqual(
                Document(output).paragraphs[1].text, "依据《新标准》GB 100-2026 执行。"
            )
            self.assertTrue(Document(output).paragraphs[1].runs[0].bold)

            diff = validate_word_change_diff(source, output, changes)
            self.assertTrue(diff["ok"])
            self.assertEqual(diff["verified_change_count"], 2)
            self.assertEqual(diff["changed_locator_count"], 1)

            tampered = Document(output)
            tampered.paragraphs[0].text = "被意外修改的标题"
            tampered.save(output)
            unsafe_diff = validate_word_change_diff(source, output, changes)
            self.assertFalse(unsafe_diff["ok"])
            self.assertIn("outside_target_change", unsafe_diff["errors"])


def make_change(
    locator: dict[str, object], target: str, after: str
) -> dict[str, object]:
    span = dict(locator[target])
    return {
        "change_id": target,
        "locator": {
            key: value
            for key, value in locator.items()
            if key not in ("standard_no", "standard_name")
        },
        "span": span,
        "before": span["text"],
        "after": after,
    }


if __name__ == "__main__":
    unittest.main()
